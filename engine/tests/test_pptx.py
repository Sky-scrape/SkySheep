"""PPT（.pptx）读写：read_document 抽取文本、write_document 生成演示稿。"""

from __future__ import annotations

import pytest
from pptx import Presentation

from skysheep.tools import ToolContext, ToolError, ToolRegistry, default_tools
from skysheep.tools.docs import (
    ReadDocumentArgs,
    WriteDocumentArgs,
    _parse_slide_markdown,
    build_pptx,
)


def _ctx(tmp_path):
    return ToolContext(tmp_path)


# ---------------------------------------------------------------- 生成


def test_build_pptx_cover_and_pages(tmp_path):
    target = tmp_path / "deck.pptx"
    content = """# 季度复盘

## 本季成果
- 上线 3 个功能
- 用户增长 12%

## 下季计划
- 补齐移动端
- 优化首屏加载
"""
    summary = build_pptx(target, content)
    assert target.exists()
    prs = Presentation(str(target))
    assert len(prs.slides) == 3  # 封面 + 2 页
    assert "季度复盘" in summary or "3 页" in summary
    texts = [
        "".join(r.text for p in s.shapes[0].text_frame.paragraphs for r in p.runs)
        for s in prs.slides
    ]
    assert "季度复盘" in texts[0]
    assert "本季成果" in texts[1]
    assert "下季计划" in texts[2]


def test_build_pptx_bullets_and_notes(tmp_path):
    target = tmp_path / "deck.pptx"
    content = """# 标题

## 要点页
- 第一条
- 第二条
  - 子项
> 这是演讲者备注
"""
    build_pptx(target, content)
    prs = Presentation(str(target))
    slide = prs.slides[1]
    body = slide.placeholders[1].text_frame
    bullet_texts = [p.text for p in body.paragraphs if p.text]
    assert "第一条" in bullet_texts and "第二条" in bullet_texts and "子项" in bullet_texts
    # 缩进的子项层级更深
    levels = {p.text: p.level for p in body.paragraphs}
    assert levels["第一条"] == 0
    assert levels["子项"] >= 1
    assert "备注" in slide.notes_slide.notes_text_frame.text


def test_build_pptx_numbered_list_and_subheading(tmp_path):
    target = tmp_path / "deck.pptx"
    content = """# T

## 步骤
### 准备阶段
1. 收集数据
2. 清洗数据
"""
    build_pptx(target, content)
    prs = Presentation(str(target))
    body = prs.slides[1].placeholders[1].text_frame
    texts = [p.text for p in body.paragraphs if p.text]
    assert "准备阶段" in texts
    assert "收集数据" in texts and "清洗数据" in texts


def test_build_pptx_requires_content(tmp_path):
    with pytest.raises(ToolError) as ei:
        build_pptx(tmp_path / "empty.pptx", "   \n  ")
    assert "没有可生成的内容" in str(ei.value)


def test_build_pptx_paragraph_without_headings(tmp_path):
    """没有标题层级的散行也要成页（不能因缺 ## 就丢内容）。"""
    target = tmp_path / "plain.pptx"
    build_pptx(target, "# 封面\n\n一段说明文字\n\n## 一页\n- 一项")
    prs = Presentation(str(target))
    all_text = " ".join(
        p.text for s in prs.slides for sh in s.shapes
        if sh.has_text_frame for p in sh.text_frame.paragraphs
    )
    assert "一段说明文字" in all_text


def test_build_pptx_warns_on_overflow(tmp_path):
    target = tmp_path / "big.pptx"
    bullets = "\n".join(f"- 条目 {i}" for i in range(20))
    content = f"# 大页\n\n## 内容\n{bullets}\n"
    summary = build_pptx(target, content)
    assert "内容较多" in summary and "建议拆分" in summary
    assert target.exists()


def test_parse_slide_markdown_structure():
    cover, slides = _parse_slide_markdown("# 封面\n\n## 页一\n- a\n> 备注一\n\n## 页二\n- b\n")
    assert cover == "封面"
    assert [s.title for s in slides] == ["页一", "页二"]
    assert slides[0].bullets == [(0, "a")]
    assert slides[0].notes == ["备注一"]


# ---------------------------------------------------------------- 工具层生成


async def test_write_document_pptx(tmp_path):
    tool = ToolRegistry(default_tools()).get("write_document")
    out = await tool.run(
        WriteDocumentArgs(
            path="out.pptx",
            content="# 演示\n\n## 第一页\n- 要点 A\n- 要点 B\n",
        ),
        _ctx(tmp_path),
    )
    assert "PPT 已生成" in out
    assert (tmp_path / "out.pptx").exists()
    assert len(Presentation(str(tmp_path / "out.pptx")).slides) == 2


async def test_write_document_pptx_rejects_empty(tmp_path):
    tool = ToolRegistry(default_tools()).get("write_document")
    with pytest.raises(ToolError):
        await tool.run(WriteDocumentArgs(path="x.pptx", content="   "), _ctx(tmp_path))


# ---------------------------------------------------------------- 读取


def _make_deck(path, lines=("标题甲", "要点一", "要点二")):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "季度汇报"
    slide.placeholders[1].text = lines[0]
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "成果"
    body = s2.placeholders[1].text_frame
    body.text = lines[1]
    body.add_paragraph().text = lines[2]
    s2.notes_slide.notes_text_frame.text = "记得强调增长数据"
    prs.save(str(path))


async def test_read_document_pptx(tmp_path):
    _make_deck(tmp_path / "d.pptx")
    tool = ToolRegistry(default_tools()).get("read_document")
    out = await tool.run(ReadDocumentArgs(path="d.pptx"), _ctx(tmp_path))
    assert "共 2 页" in out
    assert "--- 第 1 页 ---" in out and "--- 第 2 页 ---" in out
    assert "季度汇报" in out
    assert "要点一" in out and "要点二" in out
    # 备注也要抽出来（演示稿要点常写在备注里）
    assert "记得强调增长数据" in out


async def test_read_document_pptx_includes_tables(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # 只有标题的版式
    slide.shapes.title.text = "数据"
    rows, cols = 2, 2
    table = slide.shapes.add_table(rows, cols, 0, 0, 100, 100).table
    table.cell(0, 0).text = "月份"
    table.cell(0, 1).text = "销量"
    table.cell(1, 0).text = "1月"
    table.cell(1, 1).text = "120"
    prs.save(str(tmp_path / "t.pptx"))

    tool = ToolRegistry(default_tools()).get("read_document")
    out = await tool.run(ReadDocumentArgs(path="t.pptx"), _ctx(tmp_path))
    assert "| 月份 | 销量 |" in out
    assert "| 1月 | 120 |" in out


async def test_read_document_pptx_empty_page_label(tmp_path):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])  # 空白版式，无文本
    prs.save(str(tmp_path / "blank.pptx"))
    tool = ToolRegistry(default_tools()).get("read_document")
    out = await tool.run(ReadDocumentArgs(path="blank.pptx"), _ctx(tmp_path))
    assert "本页无文本" in out


async def test_read_document_supported_lists_pptx(tmp_path):
    """错误提示里要出现 pptx，用户才知道该换什么格式。"""
    (tmp_path / "x.txt").write_text("hi", encoding="utf-8")
    tool = ToolRegistry(default_tools()).get("read_document")
    with pytest.raises(ToolError) as ei:
        await tool.run(ReadDocumentArgs(path="x.txt"), _ctx(tmp_path))
    assert ".pptx" in str(ei.value)


# ---------------------------------------------------------------- 往返


async def test_pptx_roundtrip(tmp_path):
    """生成 → 读回：内容不丢（模型自己造稿再看自己的稿是最常见的用法）。"""
    reg = ToolRegistry(default_tools())
    content = "# 产品介绍\n\n## 核心能力\n- 本地运行\n- 权限确认制\n"
    await reg.get("write_document").run(
        WriteDocumentArgs(path="r.pptx", content=content), _ctx(tmp_path)
    )
    out = await reg.get("read_document").run(
        ReadDocumentArgs(path="r.pptx"), _ctx(tmp_path)
    )
    assert "产品介绍" in out
    assert "核心能力" in out
    assert "本地运行" in out and "权限确认制" in out
