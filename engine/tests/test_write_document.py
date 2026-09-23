"""write_document 文档生成工具测试：docx / xlsx / csv 生成 + 回读校验 + 检查点 + 权限。"""

from __future__ import annotations

from skysheep.security.gate import PermissionGate
from skysheep.tools import ChangeRecorder, ToolRegistry, default_tools
from skysheep.tools.base import ToolContext, ToolError
from skysheep.tools.docs import (
    ReadDocumentArgs,
    ReadDocumentTool,
    WriteDocumentArgs,
    WriteDocumentTool,
)


def make_ctx(tmp_path):
    return ToolContext(working_dir=tmp_path)


DOC_MD = """# 季度报告

这是**加粗**的一段话，带 `code` 片段。

## 数据明细

| 月份 | 销售额 |
| --- | --- |
| 1月 | 12000 |
| 2月 | 13500 |

- 第一点
- 第二点

> 结论：稳步增长。
"""


async def test_write_document_docx_roundtrip(tmp_path):
    tool = WriteDocumentTool()
    result = await tool.run(
        WriteDocumentArgs(path="报告.docx", content=DOC_MD), make_ctx(tmp_path)
    )
    assert "报告.docx" in result

    # 回读：走 read_document 同一条解析链路
    back = await ReadDocumentTool().run(
        ReadDocumentArgs(path=str(tmp_path / "报告.docx")), make_ctx(tmp_path)
    )
    assert "季度报告" in back
    assert "加粗" in back
    assert "月份" in back and "12000" in back  # 表格进了文档
    assert "第一点" in back
    assert "稳步增长" in back


async def test_write_document_docx_new_file_diff_and_checkpoint(tmp_path):
    rec = ChangeRecorder()
    tool = WriteDocumentTool(recorder=rec)
    ctx = make_ctx(tmp_path)
    await tool.run(WriteDocumentArgs(path="r.docx", content=DOC_MD), ctx)
    # 新建：改前不存在（回滚时应删除），diff 为源内容全文新增（挂在 ctx 上，
    # 不在工具实例上——实例会被并行任务共享）
    assert list(rec.pre) == [str(tmp_path / "r.docx")]
    assert rec.pre[str(tmp_path / "r.docx")] is None
    assert ctx.last_diff and "+# 季度报告" in ctx.last_diff


async def test_write_document_docx_overwrite_no_diff_and_checkpoint(tmp_path):
    p = tmp_path / "r.docx"
    p.write_bytes(b"OLDBINARY")
    rec = ChangeRecorder()
    tool = WriteDocumentTool(recorder=rec)
    ctx = make_ctx(tmp_path)
    await tool.run(WriteDocumentArgs(path=str(p), content=DOC_MD), ctx)
    assert rec.pre[str(p)] == b"OLDBINARY"  # 覆盖前快照 → 可回滚
    assert ctx.last_diff == ""  # 旧内容是二进制，给不出有意义的文本 diff


async def test_write_document_xlsx_roundtrip(tmp_path):
    content = (
        '{"销售数据": [["月份", "销售额"], ["1月", 12000], ["2月", 13500.5], '
        "[null, true]], \"备注\": [[\"说明\", \"这是备注\"]]}"
    )
    result = await WriteDocumentTool().run(
        WriteDocumentArgs(path="data.xlsx", content=content), make_ctx(tmp_path)
    )
    assert "2 个工作表" in result

    back = await ReadDocumentTool().run(
        ReadDocumentArgs(path=str(tmp_path / "data.xlsx")), make_ctx(tmp_path)
    )
    assert "销售数据" in back and "备注" in back
    assert "12000" in back and "13500.5" in back


async def test_write_document_xlsx_sheet_list_and_single_forms(tmp_path):
    # 形态二：[{name, rows}]；形态三：裸二维数组（单表 Sheet1）
    r1 = await WriteDocumentTool().run(
        WriteDocumentArgs(path="a.xlsx", content='[{"name": "表A", "rows": [["x", 1]]}]'),
        make_ctx(tmp_path),
    )
    assert "表A" in r1
    r2 = await WriteDocumentTool().run(
        WriteDocumentArgs(path="b.xlsx", content='[["h1", "h2"], ["v1", 2]]'),
        make_ctx(tmp_path),
    )
    assert "Sheet1" in r2

    import openpyxl

    wb = openpyxl.load_workbook(tmp_path / "b.xlsx")
    assert wb.sheetnames == ["Sheet1"]
    ws = wb["Sheet1"]
    assert ws.cell(1, 1).value == "h1"
    assert ws.cell(2, 2).value == 2


async def test_write_document_csv_bom_and_rows(tmp_path):
    result = await WriteDocumentTool().run(
        WriteDocumentArgs(path="t.csv", content="月份,销售额\n1月,12000\n2月,13500"),
        make_ctx(tmp_path),
    )
    assert "3 行" in result
    raw = (tmp_path / "t.csv").read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # BOM：Excel 双击打开中文不乱码
    assert b"13500" in raw


async def test_write_document_errors(tmp_path):
    tool = WriteDocumentTool()
    for exc, args in [
        ("只支持", WriteDocumentArgs(path="a.pdf", content="x")),  # 不支持的扩展名
        ("不能为空", WriteDocumentArgs(path="a.docx", content="   ")),
        ("必须是 JSON", WriteDocumentArgs(path="a.xlsx", content="不是 JSON")),
        ("没有任何工作表", WriteDocumentArgs(path="a.xlsx", content="{}")),
    ]:
        try:
            await tool.run(args, make_ctx(tmp_path))
            raised = ""
        except ToolError as e:
            raised = str(e)
        assert exc in raised, f"{args.path}: 期望报错含「{exc}」，实际：{raised}"


async def test_write_document_registered_and_gate_preview(tmp_path):
    # 注册进默认工具集（WRITE 分级），权限门对新建文档给出源内容 diff 预览
    reg = ToolRegistry(default_tools())
    assert reg.get("write_document") is not None

    gate = PermissionGate(working_dir=tmp_path)
    pending = await gate.authorize(
        WriteDocumentTool(), {"path": "新文档.docx", "content": "# 标题\n正文"}
    )
    assert pending is not None  # WRITE 需确认
    assert "+# 标题" in pending.diff

    # 覆盖已有 xlsx：无 diff；csv 是文本：真实 diff
    (tmp_path / "e.xlsx").write_bytes(b"BIN")
    p2 = await gate.authorize(
        WriteDocumentTool(), {"path": "e.xlsx", "content": "[[1]]"}
    )
    assert p2.diff == ""
    (tmp_path / "e.csv").write_text("a,b\n1,2", encoding="utf-8-sig")
    p3 = await gate.authorize(
        WriteDocumentTool(), {"path": "e.csv", "content": "a,b\n1,3"}
    )
    assert "-1,2" in p3.diff and "+1,3" in p3.diff
