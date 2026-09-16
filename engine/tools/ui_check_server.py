"""浏览器验证：起一个隔离 SKYSHEEP_HOME 的服务端，供 UI 检查脚本连。

用法（在 engine/ 目录下）：python tools/ui_check_server.py
- 应用在 http://127.0.0.1:8771/，HOME/项目目录都在系统临时目录里；
- 8772 端口挂了一个静态桩，提供 pack.zip（模拟 GitHub 下载），并临时把
  127.0.0.1 加进下载白名单——仅本开发工具生效，不影响正式行为。
- v0.7.0 用户旅程夹具：项目里预置 docx/pdf/xlsx 样例（文件树与 read_document
  演示），FakeProvider 脚本依次演练 写文件→读文档→联网搜索→记忆→画图。
"""

from __future__ import annotations

import http.server
import json
import os
import socketserver
import sys
import threading
import zipfile
from pathlib import Path

BASE = Path(os.environ.get("TEMP", "/tmp")) / "skysheep-ui-check"
HOME = BASE / "home"
PROJ = BASE / "proj"
DEMO = BASE / "demo-skill"
WWW = BASE / "www"

HOME.mkdir(parents=True, exist_ok=True)
PROJ.mkdir(parents=True, exist_ok=True)
WWW.mkdir(parents=True, exist_ok=True)

# 预置一个待导入的技能文件夹（外部目录，模拟用户从网上下载的）
DEMO.mkdir(parents=True, exist_ok=True)
(DEMO / "SKILL.md").write_text(
    "---\nname: demo-skill\ndescription: 验证导入用的示例技能\n---\n第一步：观察。\n", encoding="utf-8"
)
(DEMO / "run.py").write_text("print('demo')\n", encoding="utf-8")

# 预置一个 zip 技能包（带 repo-main 外层，模拟 GitHub 归档）
with zipfile.ZipFile(WWW / "pack.zip", "w") as zf:
    zf.writestr(
        "repo-main/skills/url-skill/SKILL.md",
        "---\nname: url-skill\ndescription: 从网址装进来的技能\n---\n内容\n",
    )

os.environ["SKYSHEEP_HOME"] = str(HOME)
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

# 预置项目文件：@ 提及 / 文件树 / read_document 演示
(PROJ / "notes.md").write_text("# 笔记\n\n这是一份验证用的笔记。\n", encoding="utf-8")
(PROJ / "hello.html").write_text(
    "<h1>预览演示页</h1><p>这是文件树一键预览的 HTML。</p>", encoding="utf-8"
)
(PROJ / "src").mkdir(exist_ok=True)
(PROJ / "src" / "demo.py").write_text("print('demo')\n", encoding="utf-8")

import docx  # noqa: E402
import openpyxl  # noqa: E402
from pypdf import PdfWriter  # noqa: E402

_d = docx.Document()
_d.add_paragraph("这是季度报告正文：SkySheep 表现优异。")
_t = _d.add_table(rows=1, cols=2)
_t.rows[0].cells[0].text = "指标"
_t.rows[0].cells[1].text = "数值"
_d.save(str(PROJ / "季度报告.docx"))

_wb = openpyxl.Workbook()
_ws = _wb.active
_ws.title = "数据"
_ws.append(["月份", "新增用户"])
_ws.append(["八月", 120])
_wb.save(str(PROJ / "数据表.xlsx"))

_pw = PdfWriter()
_page = _pw.add_blank_page(width=612, height=792)
_pw.write(str(PROJ / "说明文档.pdf"))

# 第二个项目目录，供「工作项目切换」验证
PROJ2 = BASE / "proj2"
PROJ2.mkdir(parents=True, exist_ok=True)
(PROJ2 / "README.md").write_text("# 项目二\n\n切换工作项目后的验证目录。\n", encoding="utf-8")

from skysheep.skills import installer as skill_installer  # noqa: E402

skill_installer.TRUSTED_ZIP_HOSTS = frozenset(
    set(skill_installer.TRUSTED_ZIP_HOSTS) | {"127.0.0.1", "localhost"}
)

import uvicorn  # noqa: E402

from skysheep.messages import TextBlock, ToolUseBlock  # noqa: E402
from skysheep.models.fake import FakeProvider  # noqa: E402
from skysheep.server import create_app  # noqa: E402


class _Handler(http.server.SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(WWW), **kwargs)

    def log_message(self, *args):
        pass


socketserver.TCPServer.allow_reuse_address = True
stub = socketserver.TCPServer(("127.0.0.1", 8772), _Handler)
threading.Thread(target=stub.serve_forever, daemon=True).start()

# FakeProvider：无 Key 也能跑完整对话流。脚本组按 stream 调用顺序消费，
# 用完后走 with_default。每轮「工具调用→收尾文本」消耗两组，连续排布：
provider = FakeProvider(
    [
        # 第 1 轮对话：写文件 → 权限确认 → 检查点条
        [ToolUseBlock(id="t1", name="write_file", input={"path": "ui-demo.txt", "content": "演示"})],
        [TextBlock(text="已写入 ui-demo.txt")],
        # 第 2 轮：read_document 读 docx
        [ToolUseBlock(id="t2", name="read_document", input={"path": "季度报告.docx"})],
        [TextBlock(text="读完了：季度报告正文提到 SkySheep 表现优异，还有一个指标表。")],
        # 第 3 轮：web_search（未配置 → 应给配置指引卡片而不是崩溃）
        [ToolUseBlock(id="t3", name="web_search", input={"query": "SkySheep"})],
        [TextBlock(text="搜索工具返回了配置指引。")],
        # 第 4 轮：memory_write 记住偏好
        [ToolUseBlock(
            id="t4", name="memory_write",
            input={"action": "append", "content": "用户喜欢简洁回复"},
        )],
        [TextBlock(text="好的，已记住你喜欢简洁回复。")],
        # 第 5 轮：generate_image（未配置 → 配置指引）
        [ToolUseBlock(id="t5", name="generate_image", input={"prompt": "一只云朵小羊"})],
        [TextBlock(text="画图工具返回了配置指引。")],
    ]
)
provider.with_default([TextBlock(text="收到！这是演示回复。")])

app = create_app(working_dir=str(PROJ), provider_name="fake", provider_factory=lambda: provider)
config = uvicorn.Config(app, host="127.0.0.1", port=8771, log_level="warning")
server = uvicorn.Server(config)
print(
    f"serving on http://127.0.0.1:8771/ (zip stub :8772)  HOME={HOME}  PROJ={PROJ}",
    flush=True,
)
print("demo skill:", DEMO, flush=True)
print("demo zip: http://127.0.0.1:8772/pack.zip", flush=True)
print("mcp json hint:", json.dumps({"command": sys.executable}), flush=True)
threading.Thread(target=server.run, daemon=True).start()
try:
    while not server.should_exit:
        threading.Event().wait(1)
except KeyboardInterrupt:
    pass
